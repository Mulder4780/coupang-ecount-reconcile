# -*- coding: utf-8 -*-
"""표준 업무 절차서 — 지정 폴더로 워드·PPT·유니웍스 넘김 (2026-09-10 형님 지시)

형님 지시: "절차서는 내가 지정한 폴더에 별도로 word나 ppt로 만들 수 있는 구조
준비해 … 나중에 유니웍스에 자료 넘겨서 유니웍스에서 통합관리할거야."

★ 담는 자리는 `sop_store` 하나다([162]).  여기는 **그것을 읽어 내보내기만**
  한다 — 여기서 절차서를 고치거나 만들지 않는다.  두 곳이 각자 담으면
  화면과 워드가 서로 다른 절차서를 말한다.

★ 0건이면 **파일을 안 만든다**([169]).  빈 워드를 내보내면 받아 본 사람이
  "절차서가 이게 전부구나"로 읽는다.  못 만들었으면 못 만들었다고 말한다.

★ 라이브러리가 없으면 **지어내지 않는다**([169]).  python-docx·python-pptx 가
  없으면 그 사실과 깔 명령을 말하고 멈춘다 — 반쪽 파일을 남기지 않는다.

★ Z: 를 기본 폴더로 두지 않는다([168]).  회차가 공유폴더를 물고 있으면
  내보내기 한 번이 몇 분씩 선다.
"""

import io
import json
import os
import re
import sys

if hasattr(sys.stdout, "reconfigure"):          # 무인 회차는 sys.stdout 이 None 이다([235])
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

BASE = os.path.dirname(os.path.abspath(__file__))
if BASE not in sys.path:
    sys.path.insert(0, BASE)

import sop_store as S                            # noqa: E402

# 유니웍스로 넘기는 판.  ★ 모양을 바꾸면 이 숫자를 올린다 — 받는 쪽이
# "이게 어느 판 자료냐"를 물을 수 있어야 한다([169]).
UNIWORKS_SCHEMA = 2      # 2: 칸 '그림'(사진·도면 이름·설명·단계) 추가 · 2026-09-22


def _need(mod, pip_name):
    """라이브러리를 못 찾으면 **깔 명령까지** 말한다([448] — 조치는 실재하는 명령이다)."""
    try:
        return __import__(mod)
    except ImportError:
        raise RuntimeError(
            "%s 라이브러리가 없어 만들 수 없다 — 먼저 깐다:\n"
            "  python -m pip install %s" % (mod, pip_name))


def _prep(path):
    """내보낼 자리를 준비한다.  폴더는 **여기서** 만든다(설정할 때가 아니라)."""
    d = os.path.dirname(os.path.abspath(path))
    os.makedirs(d, exist_ok=True)
    return path


def _rows(ids=None, 부위=None, 작업종류=None, store=None, 묶음=None):
    d = S.load(store)
    if d.get("_못읽음"):
        raise RuntimeError("절차서 판을 못 읽었다: " + d["_못읽음"])
    rows = S.pick(d, ids=ids, 부위=부위, 작업종류=작업종류, 묶음=묶음)
    if not rows:
        if 묶음:
            raise RuntimeError("'%s' 묶음이 없거나 비었다 — 빈 파일을 만들지 않는다.\n"
                               "  있는 묶음: %s" % (묶음, ", ".join(b["이름"] for b in S.books(d)) or "없음"))
        raise RuntimeError(
            "내보낼 절차서가 0건이다 — 빈 파일을 만들지 않는다.\n"
            "  담긴 것: %d건 (python sop_store.py 로 본다)" % len(d.get("절차서", [])))
    return rows


def _그림들(r, 단계=None):
    """그 절차서의 그림 중 `단계` 에 붙는 것(None 이면 전부)."""
    out = []
    for g in r.get("그림") or []:
        if not isinstance(g, dict):
            continue
        if 단계 is not None and int(g.get("단계") or 0) != 단계:
            continue
        out.append(g)
    return out


def _docx_그림(doc, g):
    """그림 한 장을 넣는다.  ★ 파일이 없으면 **없다고 적는다**([169]) —
    조용히 빼면 받아 본 사람은 원래 그림이 없는 장으로 읽는다."""
    from docx.shared import Cm
    p = S.그림경로(g.get("파일"))
    설명 = g.get("설명") or ""
    try:
        if not p or not os.path.isfile(p):
            raise FileNotFoundError(p)
        doc.add_picture(p, width=Cm(14))
    except Exception as e:                      # noqa: BLE001 - 깨진 그림 하나가 책 전체를 막지 않게
        doc.add_paragraph("(그림을 넣지 못함: %s — %s)" % (g.get("파일"), type(e).__name__))
    if 설명:
        cap = doc.add_paragraph("그림 — " + 설명)
        cap.runs[0].italic = True


# ── 워드 ──────────────────────────────────────────────────────────────
def export_docx(path, ids=None, 부위=None, 작업종류=None, store=None, 묶음=None):
    docx = _need("docx", "python-docx")
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    rows = _rows(ids, 부위, 작업종류, store, 묶음)
    doc = docx.Document()

    st = doc.styles["Normal"]
    st.font.name = "맑은 고딕"
    st.font.size = Pt(10.5)

    h = doc.add_heading(묶음 or "쿠팡 표준 업무 절차서", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("전 %d장 · 만든 때 %s" % (len(rows), S._now()))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if 묶음:
        # 책이면 차례를 앞에 세운다 — 신입은 무엇이 어디 있는지부터 찾는다
        doc.add_heading("차례", level=1)
        for i, r in enumerate(rows, 1):
            doc.add_paragraph("%d. %s" % (i, r.get("제목") or ""))

    for i, r in enumerate(rows):
        if i:
            doc.add_page_break()                 # 한 절차서 = 한 장(현장에서 뜯어 쓴다)
        doc.add_heading("%s  %s" % (r.get("id") or "", r.get("제목") or ""), level=1)

        # ★ 빈 칸을 '없음'이라 적지 않는다([169]) — '아직 안 적었다'와 '없다'는 다르다
        meta = [("부위", r.get("부위") or "(아직 안 적음)"),
                ("작업 종류", r.get("작업종류") or "(아직 안 적음)"),
                ("출처", r.get("출처") or "(아직 안 적음)"),
                ("작성자", r.get("작성자") or "(아직 안 적음)"),
                ("수정일", r.get("수정일") or "")]
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for k, v in meta:
            c = t.add_row().cells
            c[0].text = k
            c[1].text = str(v)

        공구 = r.get("공구") or []
        doc.add_heading("준비 공구", level=2)
        if 공구:
            for x in 공구:
                doc.add_paragraph(str(x), style="List Bullet")
        else:
            doc.add_paragraph("(아직 안 적음)")

        단계 = r.get("단계") or []
        doc.add_heading("작업 순서", level=2)
        if 단계:
            for s in 단계:
                # 번호를 글자로 적는다 — 'List Number' 는 장이 바뀌어도 번호가
                # 이어져(1~4 다음 장이 5부터) 현장에서 몇 번째인지 헷갈린다
                doc.add_paragraph("%d. %s" % (s.get("순서") or 0, s.get("내용") or ""))
                if s.get("주의"):
                    q = doc.add_paragraph("주의 — %s" % s["주의"])
                    q.paragraph_format.left_indent = Pt(24)
                for g in _그림들(r, int(s.get("순서") or 0)):
                    _docx_그림(doc, g)
        else:
            doc.add_paragraph("(아직 안 적음 — 작성자가 채워야 한다)")

        끝그림 = _그림들(r, 0)
        if 끝그림:
            doc.add_heading("사진·도면", level=2)
            for g in 끝그림:
                _docx_그림(doc, g)

        if r.get("마무리"):
            doc.add_heading("마무리", level=2)
            doc.add_paragraph(r["마무리"])
        if r.get("주의사항"):
            doc.add_heading("주의사항", level=2)
            doc.add_paragraph(r["주의사항"])

    doc.save(_prep(path))
    return path, len(rows)


# ── PPT ───────────────────────────────────────────────────────────────
def export_pptx(path, ids=None, 부위=None, 작업종류=None, store=None, 묶음=None):
    pptx = _need("pptx", "python-pptx")
    from pptx.util import Pt as PPt
    from pptx.util import Inches as PInches

    rows = _rows(ids, 부위, 작업종류, store, 묶음)
    prs = pptx.Presentation()

    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = 묶음 or "쿠팡 표준 업무 절차서"
    title.placeholders[1].text = "전 %d건 · 만든 때 %s" % (len(rows), S._now())

    body_layout = prs.slide_layouts[1]
    for r in rows:
        sl = prs.slides.add_slide(body_layout)
        sl.shapes.title.text = "%s  %s" % (r.get("id") or "", r.get("제목") or "")
        tf = sl.placeholders[1].text_frame
        tf.word_wrap = True

        머리 = "%s · %s" % (r.get("부위") or "부위 미기입",
                            r.get("작업종류") or "종류 미기입")
        tf.text = 머리

        공구 = r.get("공구") or []
        p = tf.add_paragraph()
        p.text = "공구: " + (", ".join(공구) if 공구 else "(아직 안 적음)")
        p.level = 1

        단계 = r.get("단계") or []
        if 단계:
            for s in 단계:
                p = tf.add_paragraph()
                p.text = "%d) %s" % (s.get("순서") or 0, s.get("내용") or "")
                p.level = 1
                if s.get("주의"):
                    q = tf.add_paragraph()
                    q.text = "주의 — %s" % s["주의"]
                    q.level = 2
        else:
            p = tf.add_paragraph()
            p.text = "작업 순서가 아직 안 적혔다 — 작성자가 채워야 한다"
            p.level = 1

        if r.get("마무리"):
            p = tf.add_paragraph()
            p.text = "마무리: " + r["마무리"]
            p.level = 1

        for para in tf.paragraphs:               # 현장에서 보는 장이라 글씨를 키운다
            for run in para.runs:
                run.font.size = PPt(16 if para.level < 2 else 13)

        # 사진·도면은 한 장에 하나씩 — 글과 섞으면 현장에서 안 보인다
        for g in _그림들(r):
            gs = prs.slides.add_slide(prs.slide_layouts[5])
            gs.shapes.title.text = (g.get("설명") or r.get("제목") or "")[:60]
            gp = S.그림경로(g.get("파일"))
            if gp and os.path.isfile(gp):
                try:
                    gs.shapes.add_picture(gp, PInches(0.5), PInches(1.5), height=PInches(5.5))
                    continue
                except Exception:               # noqa: BLE001 - 깨진 그림 하나가 전체를 막지 않게
                    pass
            box = gs.shapes.add_textbox(PInches(0.5), PInches(2), PInches(9), PInches(1))
            box.text_frame.text = "(그림을 넣지 못함: %s)" % g.get("파일")

    prs.save(_prep(path))
    return path, len(rows)


# ── 유니웍스 넘김 ─────────────────────────────────────────────────────
def export_uniworks(path, ids=None, 부위=None, 작업종류=None, store=None, 묶음=None):
    """유니웍스(노승용 매니저 앱)로 넘길 자료.

    ★ 우리 화면 모양이 아니라 **평평한 자료**로 낸다.  받는 쪽이 무엇을 받는지
      알 수 있게 `스키마판`·`만든곳`·`만든때`·`건수` 를 머리에 세운다([292] —
      비지 않는 것을 앞에 세운다).
    ★ 여기서 칸을 바꾸면 `UNIWORKS_SCHEMA` 를 올린다 — 안 올리면 받는 쪽이
      옛 판을 새 판으로 읽는다([169]).
    """
    rows = _rows(ids, 부위, 작업종류, store, 묶음)
    doc = {
        "스키마판": UNIWORKS_SCHEMA,
        "만든곳": "쿠팡 통합업무 자동화(CSOS)",
        "만든때": S._now(),
        "건수": len(rows),
        "부위목록": list(S.부위목록),
        "작업종류목록": list(S.작업종류목록),
        "절차서": [{k: r.get(k) for k in S.칸} for r in rows],
    }
    p = _prep(path)
    tmp = p + ".tmp"
    with io.open(tmp, "w", encoding="utf-8", newline="") as f:
        json.dump(doc, f, ensure_ascii=False, indent=2)
    os.replace(tmp, p)
    return p, len(rows)


FORMATS = {"docx": export_docx, "pptx": export_pptx, "json": export_uniworks}


def export(fmt, path=None, **kw):
    """한 곳에서 갈래를 고른다([162]).  경로를 안 주면 지정 폴더에 만든다."""
    fmt = (fmt or "").lower().lstrip(".")
    fn = FORMATS.get(fmt)
    if not fn:
        raise ValueError("모르는 갈래: %r (쓸 수 있는 것: %s)"
                         % (fmt, ", ".join(sorted(FORMATS))))
    if not path:
        stamp = S._now().replace("-", "").replace(":", "").replace(" ", "_")
        stem = re.sub(r'[\\/:*?"<>|\s]+', "_", kw.get("묶음") or "") or "쿠팡_표준업무절차서"
        path = os.path.join(S.out_dir(), "%s_%s.%s" % (stem, stamp, fmt))
    return fn(path, **kw)


def _main(argv):
    import argparse
    ap = argparse.ArgumentParser(description="표준 업무 절차서 내보내기")
    ap.add_argument("--fmt", default="docx", choices=sorted(FORMATS),
                    help="docx(워드) · pptx(PPT) · json(유니웍스 넘김)")
    ap.add_argument("--out", help="파일 경로(안 주면 지정 폴더에 만든다)")
    ap.add_argument("--부위", dest="bui", help="그 부위만")
    ap.add_argument("--종류", dest="jong", help="그 작업 종류만")
    ap.add_argument("--id", dest="ids", action="append", help="그 절차서만(여러 번)")
    ap.add_argument("--묶음", dest="book", help="그 책(업무별 한 파일 · 적힌 차례대로)")
    a = ap.parse_args(argv)

    try:
        p, n = export(a.fmt, a.out, ids=a.ids, 부위=a.bui, 작업종류=a.jong, 묶음=a.book)
    except RuntimeError as e:
        print("★", e)
        return 2
    print("만들었다: %s (%d건)" % (p, n))
    return 0


if __name__ == "__main__":
    sys.exit(_main(sys.argv[1:]))
